#!/usr/bin/env python3
"""Refresh pipeline/state/french-deals-5-years.json from an M&A Monitor
Bid Premia export (the "French public data.xlsx" download, filtered to France).

The state file is the five-year archive of COMPLETED / FAILED French public
offers used by the admin and client sites. Pending deals are skipped: live
deals come from the AMF offer-period list, enriched from the European tracker.

Existing deals are kept (so the window never shrinks unexpectedly) and are
refreshed from the export when their deal number is present in it.

Usage:
    python pipeline/import_french_deals.py --xlsx "French public data.xlsx"
    python pipeline/import_french_deals.py --xlsx export.xlsx --since 2021-06-22
"""
import argparse
import datetime as dt
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "state" / "french-deals-5-years.json"
ARCHIVE_STATUSES = {"Completed", "Failed", "Lapsed", "Withdrawn"}


def as_date(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip() if value else ""


def as_number(value):
    try:
        return round(float(str(value).replace(",", "")), 3)
    except (TypeError, ValueError):
        return None


def as_int(value):
    try:
        return int(float(str(value).replace(",", "")))
    except (TypeError, ValueError):
        return None


def read_export(path):
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header_index = next(
        i for i, row in enumerate(rows[:30]) if row and row[0] == "Deal Number"
    )
    header = {name: i for i, name in enumerate(rows[header_index]) if name}
    col = lambda row, name: row[header[name]] if name in header else None  # noqa: E731

    deals = []
    for row in rows[header_index + 1:]:
        if not row or not row[0]:
            continue
        deals.append({
            "deal_number": str(col(row, "Deal Number")).strip(),
            "date_announced": as_date(col(row, "Date announced")),
            "date_completed": as_date(col(row, "Date completed")),
            "target": str(col(row, "Target Name") or "").strip(),
            "target_country": col(row, "Country"),
            "target_activities": col(row, "Target Activities"),
            "industrial_sector": " ".join(str(col(row, "Industrial Sector") or "").split()),
            "ultimate_bidder": col(row, "Ultimate Bidder"),
            "bidder_country": col(row, "Ultimate Bidders country"),
            "type_of_deal": " / ".join(
                part.strip() for part in str(col(row, "Type of deal") or "").split(", ") if part.strip()
            ),
            "status_of_deal": col(row, "Status of deal"),
            "deal_attitude": col(row, "Deal Attitude"),
            "consideration_type": col(row, "Consideration type"),
            "implied_total_equity_value": as_number(col(row, "Implied Total Equity Value")),
            "disclosure_date_added": as_date(col(row, "Disclosure Date Added")),
            "disclosure_date_removed": as_date(col(row, "Disclosure Date Removed")),
            "offeror_named": "",
            "target_ordinary_shares": as_int(col(row, "Number of Target Ordinary shares")),
        })
    return deals


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, help="M&A Monitor bid premia export (France)")
    parser.add_argument("--source-label", default="", help="Source path to record in the JSON")
    parser.add_argument("--since", default="", help="Earliest announcement date (default: keep current window)")
    parser.add_argument("--out", default=str(OUT))
    args = parser.parse_args()

    out = Path(args.out)
    existing = json.loads(out.read_text(encoding="utf-8")) if out.exists() else {"deals": []}
    by_number = {d["deal_number"]: d for d in existing.get("deals", [])}
    since = args.since or min(
        (d["date_announced"] for d in by_number.values() if d.get("date_announced")),
        default=(dt.date.today() - dt.timedelta(days=5 * 365)).isoformat(),
    )

    added, refreshed = [], 0
    export = sorted(read_export(args.xlsx), key=lambda d: (d["date_announced"], d["deal_number"]))
    for deal in export:
        if deal["target_country"] != "France" or deal["date_announced"] < since:
            continue
        if deal["status_of_deal"] not in ARCHIVE_STATUSES:
            continue
        if deal["deal_number"] in by_number:
            if by_number[deal["deal_number"]] != deal:
                refreshed += 1
        else:
            added.append(deal["target"])
        by_number[deal["deal_number"]] = deal

    # Keep the existing order and append new deals in announcement order.
    deals = list(by_number.values())
    payload = {
        "source": args.source_label or str(args.xlsx),
        "sheet": "Sheet 1",
        "count": len(deals),
        "updated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
        "deals": deals,
    }
    out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"since": since, "count": len(deals), "added": added, "refreshed": refreshed}, indent=2))


if __name__ == "__main__":
    main()
